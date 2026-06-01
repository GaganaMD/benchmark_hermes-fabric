#!/usr/bin/env python3
from pathlib import Path
from html import escape
from pptx import Presentation

PPTX = Path("/Users/basethesis/Desktop/benchmark_hermes-fabric/DD/Vlayx_DD_work/output/Vlayx_Due_Diligence_Report.pptx")
HTML = Path("/Users/basethesis/Desktop/benchmark_hermes-fabric/DD/Vlayx_DD_work/output/Vlayx_Due_Diligence_Report_preview.html")


def shape_text(shape):
    chunks = []
    if hasattr(shape, "text") and shape.text.strip():
        chunks.append(shape.text.strip())
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        chunks.append(rows)
    return chunks


def main():
    prs = Presentation(PPTX)
    parts = [
        "<!doctype html><meta charset='utf-8'>",
        "<title>Vlayx Due Diligence Report Preview</title>",
        "<style>",
        "body{margin:0;background:#e9edf2;font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;color:#26313f}",
        ".deck{max-width:1180px;margin:24px auto;padding:0 16px}",
        ".slide{background:white;border:1px solid #cfd6df;border-radius:8px;box-shadow:0 4px 18px rgba(24,35,52,.12);margin:0 0 24px;padding:30px;min-height:560px}",
        "h1{font-size:28px;margin:0 0 18px;color:#172334}.shape{white-space:pre-wrap;font-size:15px;line-height:1.35;margin:10px 0}",
        "table{border-collapse:collapse;width:100%;margin:12px 0;font-size:13px}td,th{border:1px solid #ccd3dc;padding:7px;vertical-align:top}tr:first-child td{background:#24446c;color:white;font-weight:600}",
        ".num{font-size:12px;color:#687385;margin-bottom:8px}",
        "</style><div class='deck'>",
    ]
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"<section class='slide'><div class='num'>Slide {i}</div>")
        first = True
        for shape in slide.shapes:
            for item in shape_text(shape):
                if isinstance(item, list):
                    parts.append("<table>")
                    for row in item:
                        parts.append("<tr>" + "".join(f"<td>{escape(cell)}</td>" for cell in row) + "</tr>")
                    parts.append("</table>")
                elif first and item:
                    parts.append(f"<h1>{escape(item)}</h1>")
                    first = False
                else:
                    parts.append(f"<div class='shape'>{escape(item)}</div>")
        parts.append("</section>")
    parts.append("</div>")
    HTML.write_text("\n".join(parts), encoding="utf-8")
    print(HTML)


if __name__ == "__main__":
    main()

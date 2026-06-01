#!/usr/bin/env python3
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import fitz
import openpyxl
import xlrd
from pptx import Presentation


ROOT = Path("/Users/basethesis/Desktop/benchmark_hermes-fabric/DD/Vlayx inputs - for testing")
WORK = Path("/Users/basethesis/Desktop/benchmark_hermes-fabric/DD/Vlayx_DD_work")
TEXT_DIR = WORK / "extracted_text"
ZIP_DIR = WORK / "zip_expanded"
OUT = WORK / "output"


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def safe_name(s: str) -> str:
    h = hashlib.sha1(s.encode("utf-8", "ignore")).hexdigest()[:10]
    base = re.sub(r"[^A-Za-z0-9_.-]+", "_", s)[-120:]
    return f"{base}.{h}.txt"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def trim_cell(v):
    if v is None:
        return ""
    s = str(v).replace("\r", " ").replace("\n", " ").strip()
    return s


def extract_xlsx(path: Path) -> str:
    out = []
    load_path = path
    temp_path = None
    if path.suffix.lower() != ".xlsx":
        fd, temp_name = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        temp_path = Path(temp_name)
        shutil.copyfile(path, temp_path)
        load_path = temp_path
    wb = openpyxl.load_workbook(load_path, data_only=False, read_only=True)
    out.append(f"Workbook sheets: {', '.join(wb.sheetnames)}")
    for ws in wb.worksheets:
        out.append(f"\n--- Sheet: {ws.title} ({ws.max_row} rows x {ws.max_column} cols) ---")
        rows_written = 0
        for row in ws.iter_rows(values_only=True):
            vals = [trim_cell(v) for v in row]
            if any(vals):
                out.append("\t".join(vals))
                rows_written += 1
            if rows_written >= 2500:
                out.append("[TRUNCATED_AFTER_2500_NONEMPTY_ROWS]")
                break
    if temp_path:
        temp_path.unlink(missing_ok=True)
    return "\n".join(out)


def extract_xls(path: Path) -> str:
    out = []
    book = xlrd.open_workbook(path, on_demand=True)
    out.append(f"Workbook sheets: {', '.join(book.sheet_names())}")
    for sheet_name in book.sheet_names():
        sh = book.sheet_by_name(sheet_name)
        out.append(f"\n--- Sheet: {sheet_name} ({sh.nrows} rows x {sh.ncols} cols) ---")
        rows_written = 0
        for r in range(sh.nrows):
            vals = [trim_cell(sh.cell_value(r, c)) for c in range(sh.ncols)]
            if any(vals):
                out.append("\t".join(vals))
                rows_written += 1
            if rows_written >= 2500:
                out.append("[TRUNCATED_AFTER_2500_NONEMPTY_ROWS]")
                break
    return "\n".join(out)


def extract_pdf(path: Path) -> str:
    doc = fitz.open(path)
    out = [f"PDF pages: {doc.page_count}"]
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        out.append(f"\n--- Page {i} ---")
        out.append(text.strip() if text.strip() else "[NO_EXTRACTABLE_TEXT_ON_PAGE]")
    return "\n".join(out)


def extract_pptx(path: Path) -> str:
    prs = Presentation(path)
    out = [f"Slides: {len(prs.slides)}"]
    for idx, slide in enumerate(prs.slides, start=1):
        out.append(f"\n--- Slide {idx} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                out.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(out)


def extract_docx(path: Path) -> str:
    out = []
    with zipfile.ZipFile(path) as zf:
        for name in ("word/document.xml",):
            if name in zf.namelist():
                xml = zf.read(name).decode("utf-8", "ignore")
                xml = re.sub(r"</w:p[^>]*>", "\n", xml)
                xml = re.sub(r"<[^>]+>", " ", xml)
                out.append(re.sub(r"\s+", " ", xml).strip())
    return "\n".join(out)


def extract_text(path: Path) -> str:
    for enc in ("utf-8", "utf-16", "utf-16-le", "utf-16-be", "latin-1"):
        try:
            return path.read_text(encoding=enc, errors="strict")
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="latin-1", errors="ignore")


def extract_csv(path: Path) -> str:
    txt = extract_text(path)
    try:
        rows = list(csv.reader(txt.splitlines()))
        return "\n".join("\t".join(row) for row in rows)
    except Exception:
        return txt


def process_file(path: Path, source_kind="tree"):
    ext = path.suffix.lower()
    item = {
        "source_kind": source_kind,
        "path": str(path),
        "relative_path": rel(path),
        "extension": ext,
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "readable": True,
        "extractor": "",
        "text_path": "",
        "chars": 0,
        "status": "ok",
        "notes": "",
    }
    try:
        if path.name == ".DS_Store" or path.name.startswith("._") or "__MACOSX" in path.parts:
            text, item["extractor"], item["notes"] = "", "skip", "macOS metadata file; no substantive content"
        elif ext == ".pdf":
            text, item["extractor"] = extract_pdf(path), "pymupdf"
        elif ext == ".xlsx":
            text, item["extractor"] = extract_xlsx(path), "openpyxl"
        elif ext == ".xls":
            if path.read_bytes()[:2] == b"PK":
                text, item["extractor"] = extract_xlsx(path), "openpyxl-mislabeled-xls"
            else:
                text, item["extractor"] = extract_xls(path), "xlrd"
        elif ext == ".pptx":
            text, item["extractor"] = extract_pptx(path), "python-pptx"
        elif ext == ".docx":
            text, item["extractor"] = extract_docx(path), "docx-xml"
        elif ext == ".csv":
            text, item["extractor"] = extract_csv(path), "csv"
        elif ext == ".txt":
            text, item["extractor"] = extract_text(path), "text"
        elif ext == ".zip":
            text, item["extractor"] = zip_inventory(path), "zip-list"
        elif ext in (".jpg", ".jpeg", ".png"):
            text, item["extractor"], item["notes"] = image_metadata(path), "pillow-metadata", "image content requires visual review/OCR"
        else:
            text = extract_text(path)
            item["extractor"] = "text/binary-fallback"
            if not re.search(r"[A-Za-z0-9]", text):
                item["notes"] = "no substantive text found by fallback extractor"
    except Exception as e:
        item["readable"] = False
        item["status"] = "error"
        item["notes"] = repr(e)
        text = ""
    item["chars"] = len(text)
    out_name = safe_name(item["relative_path"])
    item["text_path"] = str(TEXT_DIR / out_name)
    (TEXT_DIR / out_name).write_text(text, encoding="utf-8", errors="ignore")
    return item


def image_metadata(path: Path) -> str:
    from PIL import Image
    with Image.open(path) as im:
        return f"Image format: {im.format}\nDimensions: {im.width} x {im.height}\nMode: {im.mode}\n"


def zip_inventory(path: Path) -> str:
    lines = []
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            lines.append(f"{info.filename}\t{info.file_size} bytes")
    return "\n".join(lines)


def expand_zips(tree_files):
    expanded = []
    for path in tree_files:
        if path.suffix.lower() != ".zip":
            continue
        dest = ZIP_DIR / safe_name(rel(path)).replace(".txt", "")
        if dest.exists():
            shutil.rmtree(dest)
        dest.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(path) as zf:
            zf.extractall(dest)
        for child in sorted(p for p in dest.rglob("*") if p.is_file()):
            expanded.append(child)
    return expanded


def main():
    TEXT_DIR.mkdir(parents=True, exist_ok=True)
    ZIP_DIR.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    for p in TEXT_DIR.glob("*.txt"):
        p.unlink()

    tree_files = sorted(p for p in ROOT.rglob("*") if p.is_file())
    expanded_files = expand_zips(tree_files)
    manifest = []
    for i, path in enumerate(tree_files + expanded_files, start=1):
        manifest.append(process_file(path, "tree" if path in tree_files else "zip-expanded"))
        if i % 50 == 0:
            print(f"processed {i}/{len(tree_files) + len(expanded_files)}", file=sys.stderr)

    with (OUT / "manifest.json").open("w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    with (OUT / "manifest.csv").open("w", newline="", encoding="utf-8") as f:
        fields = list(manifest[0].keys())
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(manifest)

    rows = []
    for m in manifest:
        txt = Path(m["text_path"]).read_text(encoding="utf-8", errors="ignore")
        words = len(re.findall(r"\w+", txt))
        rows.append({
            "relative_path": m["relative_path"],
            "extension": m["extension"],
            "status": m["status"],
            "extractor": m["extractor"],
            "chars": m["chars"],
            "words": words,
            "notes": m["notes"],
        })
    with (OUT / "coverage.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print(json.dumps({
        "tree_files": len(tree_files),
        "zip_expanded_files": len(expanded_files),
        "manifest_items": len(manifest),
        "errors": sum(1 for m in manifest if m["status"] != "ok"),
        "output": str(OUT),
    }, indent=2))


if __name__ == "__main__":
    main()
